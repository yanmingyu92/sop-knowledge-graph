"""
content_extractor.py
--------------------
Phase 1: Content Extraction Pipeline

Extracts structured content from SOP documents (.docx and .pdf) into a
Unified Document Model (UDM) containing:
  - Metadata (doc_id, title, version, owner, etc.)
  - Section hierarchy with nesting
  - Procedure steps with role assignments
  - Tables with structured data
  - Templates for programmatic filling
  - Cross-references to other documents

Usage:
    udm = extract_document("path/to/SOP-001.docx")
    save_udm(udm, "output/udm/SOP-001.json")
"""

from __future__ import annotations

import json
import logging
import re
from pathlib import Path
from typing import Any

log = logging.getLogger(__name__)

# -- Document ID regex ----------------------------------------------------
# Matches document identifiers with optional organizational prefix.
# Adapts to naming conventions such as:
#   PREFIX-S004.ER102, PREFIX-G002.1, EPI-SOP-010.ER2, GCD-SOP-204.00
# The optional organizational prefix (e.g., "ORG ") is NOT captured.
DOC_ID_RE = re.compile(
    r"(?:[A-Z]+\s+)?"
    r"((?:[A-Z]+-[SG]|EPI-SOP-|GCD-SOP[ -]|GDMS-SOP)\d{3}[\d.]*(?:[. -]?ER\d+)?)"
)


# =========================================================================
# Unified Document Model (UDM) Schema
# =========================================================================

def create_udm() -> dict[str, Any]:
    """Create an empty Unified Document Model structure."""
    return {
        "metadata": {},
        "sections": [],
        "procedures": [],
        "tables": [],
        "templates": [],
        "cross_references": [],
    }


# =========================================================================
# Document Parser Factory
# =========================================================================

def extract_document(file_path: str | Path) -> dict[str, Any]:
    """
    Main entry point: auto-detect format and extract content.

    Args:
        file_path: Path to .docx or .pdf document

    Returns:
        Unified Document Model (UDM) dictionary
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Document not found: {path}")

    suffix = path.suffix.lower()

    if suffix == ".docx":
        parser = DocxParser(path)
    elif suffix == ".pdf":
        parser = PdfParser(path)
    else:
        raise ValueError(f"Unsupported file format: {suffix}")

    log.info("Extracting content from %s (%s)", path.name, suffix)
    udm = parser.extract()
    log.info("Extraction complete: %d sections, %d procedures, %d tables",
             len(udm["sections"]), len(udm["procedures"]), len(udm["tables"]))
    return udm


def save_udm(udm: dict[str, Any], output_path: str | Path) -> None:
    """Save UDM to JSON file."""
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(udm, f, indent=2, ensure_ascii=False)
    log.info("UDM saved to %s", path)


# =========================================================================
# Base Parser Class
# =========================================================================

class BaseParser:
    """Base class for document parsers."""

    def __init__(self, file_path: Path):
        self.file_path = file_path
        self.udm = create_udm()

    def extract(self) -> dict[str, Any]:
        """Extract all content into UDM."""
        self.extract_metadata()
        self.extract_sections()
        self.extract_procedures()
        self.extract_tables()
        self.extract_templates()
        self.extract_cross_references()
        return self.udm

    def extract_metadata(self) -> None:
        raise NotImplementedError

    def extract_sections(self) -> None:
        raise NotImplementedError

    def extract_procedures(self) -> None:
        raise NotImplementedError

    def extract_tables(self) -> None:
        raise NotImplementedError

    def extract_templates(self) -> None:
        raise NotImplementedError

    def extract_cross_references(self) -> None:
        raise NotImplementedError


# =========================================================================
# .docx Parser
# =========================================================================

class DocxParser(BaseParser):
    """Parser for Microsoft Word .docx files using python-docx."""

    def __init__(self, file_path: Path):
        super().__init__(file_path)
        try:
            from docx import Document
            self.doc = Document(str(file_path))
        except ImportError as e:
            raise ImportError("python-docx not installed. Run: pip install python-docx") from e

    def extract_metadata(self) -> None:
        """Extract metadata from .docx core properties and filename."""
        core_props = self.doc.core_properties

        filename = self.file_path.stem
        doc_id_match = DOC_ID_RE.search(filename)
        doc_id = doc_id_match.group(1) if doc_id_match else filename[:30]

        # Extract title (remove doc_id and any org prefix from filename)
        title = filename
        if doc_id_match:
            title = filename[doc_id_match.end():].strip(" -_")

        self.udm["metadata"] = {
            "doc_id": doc_id,
            "title": title or core_props.title or filename,
            "type": self._infer_document_type(filename),
            "version": str(core_props.revision or ""),
            "effective_date": str(core_props.modified) if core_props.modified else "",
            "owner": core_props.author or "",
            "revision_count": int(core_props.revision or 0),
            "file_path": str(self.file_path),
            "source_format": "docx",
        }

    def _infer_document_type(self, filename: str) -> str:
        """Infer document type from filename patterns."""
        filename_lower = filename.lower()
        if "work instruction" in filename_lower or ".er" in filename_lower and "work" in filename_lower:
            return "Work Instruction"
        elif "job aid" in filename_lower or "jobaid" in filename_lower:
            return "Job Aid"
        elif "template" in filename_lower:
            return "Template"
        elif "guidance" in filename_lower or "guideline" in filename_lower:
            return "Guidance"
        elif ".g" in filename_lower or "b-g" in filename_lower:
            return "Guidance"
        elif ".s" in filename_lower or "b-s" in filename_lower or "sop" in filename_lower:
            return "SOP"
        else:
            return "Document"

    def extract_sections(self) -> None:
        """Extract section hierarchy from paragraph styles."""
        sections = []
        current_hierarchy = []

        for para in self.doc.paragraphs:
            style_name = para.style.name if para.style else ""
            text = para.text.strip()

            if not text:
                continue

            level = self._get_heading_level(style_name, text)

            if level > 0:
                section = {
                    "id": str(len(sections) + 1),
                    "level": level,
                    "title": text,
                    "content": "",
                    "subsections": [],
                }

                while current_hierarchy and current_hierarchy[-1]["level"] >= level:
                    current_hierarchy.pop()

                if current_hierarchy:
                    current_hierarchy[-1]["subsections"].append(section)
                else:
                    sections.append(section)

                current_hierarchy.append(section)

            else:
                if current_hierarchy:
                    if current_hierarchy[-1]["content"]:
                        current_hierarchy[-1]["content"] += "\n\n"
                    current_hierarchy[-1]["content"] += text

        self.udm["sections"] = sections

    def _get_heading_level(self, style_name: str, text: str) -> int:
        """Determine heading level from style name or text pattern."""
        if "Heading 1" in style_name or "Heading1" in style_name:
            return 1
        elif "Heading 2" in style_name or "Heading2" in style_name:
            return 2
        elif "Heading 3" in style_name or "Heading3" in style_name:
            return 3
        elif "Heading 4" in style_name or "Heading4" in style_name:
            return 4
        elif "Heading 5" in style_name or "Heading5" in style_name:
            return 5
        elif "Heading 6" in style_name or "Heading6" in style_name:
            return 6

        if re.match(r"^\d+\.?\s+[A-Z]", text):
            return 1
        elif re.match(r"^\d+\.\d+\.?\s+", text):
            return 2

        return 0

    def extract_procedures(self) -> None:
        """Extract procedure steps from numbered/bulleted lists."""
        procedures = []

        for para in self.doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            step_match = re.match(r"^(\d+\.(?:\d+\.)*|\w+\.)\s+(.+)", text)
            if step_match:
                step_id = step_match.group(1).rstrip(".")
                action = step_match.group(2)

                role = ""
                if ":" in action:
                    parts = action.split(":", 1)
                    if len(parts[0].split()) <= 4:
                        role = parts[0].strip()
                        action = parts[1].strip()

                procedures.append({
                    "step_id": step_id,
                    "role": role,
                    "action": action,
                    "timing": "",
                    "substeps": [],
                })

        self.udm["procedures"] = procedures

    def extract_tables(self) -> None:
        """Extract tables with structured data."""
        tables = []

        for idx, table in enumerate(self.doc.tables, 1):
            if len(table.rows) == 0:
                continue

            headers = [cell.text.strip() for cell in table.rows[0].cells]

            rows = []
            for row in table.rows[1:]:
                row_data = [cell.text.strip() for cell in row.cells]
                rows.append(dict(zip(headers, row_data)))

            table_type = self._infer_table_type(headers, rows)

            tables.append({
                "id": f"table_{idx}",
                "title": "",
                "type": table_type,
                "columns": headers,
                "rows": rows,
            })

        self.udm["tables"] = tables

    def _infer_table_type(self, headers: list[str], rows: list[dict]) -> str:
        """Infer table type from headers and content."""
        headers_lower = [h.lower() for h in headers]

        if "complete" in headers_lower or "checkbox" in headers_lower:
            return "checklist"
        elif "who" in headers_lower or "role" in headers_lower:
            return "role_assignment"
        elif "reference" in headers_lower or "document" in headers_lower:
            return "reference_table"
        else:
            return "data_table"

    def extract_templates(self) -> None:
        """Identify and extract templates."""
        templates = []

        for section in self.udm["sections"]:
            title_lower = section["title"].lower()
            if "template" in title_lower or "appendix" in title_lower:
                content = section["content"]
                fields = re.findall(r"\{([^}]+)\}", content)

                if fields:
                    templates.append({
                        "id": f"template_{section['id']}",
                        "type": self._infer_template_type(section["title"]),
                        "title": section["title"],
                        "fields": {field: f"{{{field}}}" for field in fields},
                        "content": content,
                    })

        self.udm["templates"] = templates

    def _infer_template_type(self, title: str) -> str:
        """Infer template type from title."""
        title_lower = title.lower()
        if "email" in title_lower:
            return "email"
        elif "form" in title_lower:
            return "form"
        elif "checklist" in title_lower:
            return "checklist"
        else:
            return "generic"

    def extract_cross_references(self) -> None:
        """Extract cross-references to other documents."""
        cross_refs = []
        doc_id_pattern = DOC_ID_RE

        for para in self.doc.paragraphs:
            text = para.text
            matches = doc_id_pattern.findall(text)

            for match in matches:
                if match != self.udm["metadata"]["doc_id"]:
                    cross_refs.append({
                        "source_section": "",
                        "target_doc": match,
                        "reference_type": "mention",
                    })

        for rel in self.doc.part.rels.values():
            if "hyperlink" in rel.reltype:
                target = rel.target_ref
                doc_match = doc_id_pattern.search(target)
                if doc_match:
                    cross_refs.append({
                        "source_section": "",
                        "target_doc": doc_match.group(1),
                        "reference_type": "hyperlink",
                    })

        seen = set()
        unique_refs = []
        for ref in cross_refs:
            key = (ref["target_doc"], ref["reference_type"])
            if key not in seen:
                seen.add(key)
                unique_refs.append(ref)

        self.udm["cross_references"] = unique_refs


# =========================================================================
# .pdf Parser
# =========================================================================

class PdfParser(BaseParser):
    """Parser for PDF files using PyMuPDF + pdfplumber."""

    def __init__(self, file_path: Path):
        super().__init__(file_path)
        try:
            import pymupdf
            self.pdf = pymupdf.open(str(file_path))
        except ImportError as e:
            raise ImportError("PyMuPDF not installed. Run: pip install PyMuPDF") from e

    def extract_metadata(self) -> None:
        """Extract metadata from PDF properties and filename."""
        metadata = self.pdf.metadata

        filename = self.file_path.stem
        doc_id_match = DOC_ID_RE.search(filename)
        doc_id = doc_id_match.group(1) if doc_id_match else filename[:30]

        title = filename
        if doc_id_match:
            title = filename[doc_id_match.end():].strip(" -_")

        self.udm["metadata"] = {
            "doc_id": doc_id,
            "title": title or metadata.get("title", filename),
            "type": self._infer_document_type(filename),
            "version": metadata.get("subject", ""),
            "effective_date": metadata.get("modDate", ""),
            "owner": metadata.get("author", ""),
            "revision_count": 0,
            "file_path": str(self.file_path),
            "source_format": "pdf",
        }

    def _infer_document_type(self, filename: str) -> str:
        filename_lower = filename.lower()
        if "work instruction" in filename_lower or ".er" in filename_lower and "work" in filename_lower:
            return "Work Instruction"
        elif "job aid" in filename_lower:
            return "Job Aid"
        elif "template" in filename_lower:
            return "Template"
        elif "guidance" in filename_lower:
            return "Guidance"
        elif ".g" in filename_lower or "b-g" in filename_lower:
            return "Guidance"
        else:
            return "SOP"

    def extract_sections(self) -> None:
        """Extract sections from PDF text."""
        sections = []
        full_text = ""
        for page in self.pdf:
            full_text += page.get_text()

        lines = full_text.split("\n")
        current_section = None

        for line in lines:
            line = line.strip()
            if not line:
                continue

            if re.match(r"^\d+\.?\s+[A-Z]", line) or line.isupper() and len(line) < 100:
                if current_section:
                    sections.append(current_section)
                current_section = {
                    "id": str(len(sections) + 1),
                    "level": 1,
                    "title": line,
                    "content": "",
                    "subsections": [],
                }
            elif current_section:
                if current_section["content"]:
                    current_section["content"] += "\n"
                current_section["content"] += line

        if current_section:
            sections.append(current_section)

        self.udm["sections"] = sections

    def extract_procedures(self) -> None:
        """Extract procedures from PDF text."""
        procedures = []

        for page in self.pdf:
            text = page.get_text()
            for line in text.split("\n"):
                line = line.strip()
                step_match = re.match(r"^(\d+\.(?:\d+\.)*|\w+\.)\s+(.+)", line)
                if step_match:
                    procedures.append({
                        "step_id": step_match.group(1).rstrip("."),
                        "role": "",
                        "action": step_match.group(2),
                        "timing": "",
                        "substeps": [],
                    })

        self.udm["procedures"] = procedures

    def extract_tables(self) -> None:
        """Extract tables using pdfplumber."""
        try:
            import pdfplumber
        except ImportError:
            log.warning("pdfplumber not installed - skipping table extraction")
            return

        tables = []

        with pdfplumber.open(str(self.file_path)) as pdf:
            for page_num, page in enumerate(pdf.pages, 1):
                page_tables = page.extract_tables()
                for idx, table_data in enumerate(page_tables, 1):
                    if not table_data or len(table_data) < 2:
                        continue

                    headers = table_data[0]
                    rows = [dict(zip(headers, row)) for row in table_data[1:]]

                    tables.append({
                        "id": f"table_p{page_num}_{idx}",
                        "title": "",
                        "type": "data_table",
                        "columns": headers,
                        "rows": rows,
                    })

        self.udm["tables"] = tables

    def extract_templates(self) -> None:
        """Extract templates from PDF."""
        templates = []
        for section in self.udm["sections"]:
            if "template" in section["title"].lower():
                templates.append({
                    "id": f"template_{section['id']}",
                    "type": "generic",
                    "title": section["title"],
                    "fields": {},
                    "content": section["content"],
                })
        self.udm["templates"] = templates

    def extract_cross_references(self) -> None:
        """Extract cross-references from PDF."""
        cross_refs = []
        doc_id_pattern = re.compile(r"((?:[A-Z]+-[SG]|EPI-SOP)-[\d\.]+(?:ER\d+)?)")

        for page in self.pdf:
            text = page.get_text()
            matches = doc_id_pattern.findall(text)
            for match in matches:
                if match != self.udm["metadata"]["doc_id"]:
                    cross_refs.append({
                        "source_section": "",
                        "target_doc": match,
                        "reference_type": "mention",
                    })

        seen = set()
        unique_refs = []
        for ref in cross_refs:
            key = (ref["target_doc"], ref["reference_type"])
            if key not in seen:
                seen.add(key)
                unique_refs.append(ref)

        self.udm["cross_references"] = unique_refs

    def __del__(self):
        """Close PDF on cleanup."""
        if hasattr(self, "pdf"):
            self.pdf.close()
